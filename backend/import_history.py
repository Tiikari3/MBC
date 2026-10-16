from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from sqlmodel import Session, create_engine, select

from main import User, Song, Vote


# ============================================================
# CONFIGURATION — EDIT THIS SECTION BEFORE RUNNING
# ============================================================

BASE_DIR = Path(__file__).resolve().parent

EXCEL_FILE = Path("C:/Users/logan/Documents/Music_Book_Club.xlsx")
DATABASE_FILE = BASE_DIR / "mbc_local.db"

# Spreadsheet member name -> username in your local website database.
# Replace each placeholder with the exact username you created.
USER_MAPPING = {
    "Logan": "loganworsdell",
    "Parker": "Parker",
    "Giacomo": "PapiPetro",
    "Rosey": "rosey0316",
    "TTAC": "typiont814",
    "Weston": "WestonMiller",
    "Flocco": "Flocco",
    "Jake": "awesomegear3",
}

WORKSHEET_NAME = "Rankings"

# Excel columns containing each member's rating.
RATING_COLUMNS = {
    "Logan": 11,
    "Parker": 12,
    "Giacomo": 13,
    "Rosey": 14,
    "TTAC": 15,
    "Weston": 16,
    "Flocco": 17,
    "Jake": 18,
}


# ============================================================
# DATABASE CONNECTION
# Explicitly target the local database, not Render.
# ============================================================

engine = create_engine(
    f"sqlite:///{DATABASE_FILE.as_posix()}",
    connect_args={"check_same_thread": False},
)


def clean_text(value):
    """Convert a spreadsheet cell to clean text."""
    if value is None:
        return ""

    return str(value).strip()


def parse_date(value):
    """Convert a spreadsheet date to a Python date."""
    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    if value is None or not clean_text(value):
        return None

    text_value = clean_text(value)

    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y"):
        try:
            return datetime.strptime(text_value, fmt).date()
        except ValueError:
            continue

    raise ValueError(f"Unrecognized date: {value!r}")


def parse_rating(value):
    """Return a valid whole-number rating from 1 to 7, or None."""
    if value is None or not clean_text(value):
        return None

    try:
        number = float(value)
    except (TypeError, ValueError):
        return None

    if not number.is_integer():
        return None

    score = int(number)

    if 1 <= score <= 7:
        return score

    return None


def song_key(title, artist, submitted_date, submitter_id):
    """Identify a song submission for duplicate checking."""
    return (
        title.strip().casefold(),
        artist.strip().casefold(),
        submitted_date,
        submitter_id,
    )


def main():
    # Basic safety checks.
    if not EXCEL_FILE.is_file():
        raise FileNotFoundError(
            f"Excel workbook not found: {EXCEL_FILE}"
        )

    if not DATABASE_FILE.is_file():
        raise FileNotFoundError(
            f"Local database not found: {DATABASE_FILE}\n"
            "Create your local database and accounts first."
        )

    placeholders = [
        name
        for name, username in USER_MAPPING.items()
        if not username.strip()
        or username.startswith("CHANGE_TO_")
    ]

    if placeholders:
        raise ValueError(
            "Fill in the username mapping for: "
            + ", ".join(placeholders)
        )

    workbook = load_workbook(
        EXCEL_FILE,
        read_only=True,
        data_only=True,
    )

    try:
        if WORKSHEET_NAME not in workbook.sheetnames:
            raise ValueError(
                f"Worksheet '{WORKSHEET_NAME}' was not found. "
                f"Available sheets: {workbook.sheetnames}"
            )

        sheet = workbook[WORKSHEET_NAME]

        with Session(engine) as session:
            # Verify that every mapped local account exists.
            users = session.exec(select(User)).all()
            users_by_username = {
                user.username: user for user in users
            }

            missing_users = [
                username
                for username in USER_MAPPING.values()
                if username not in users_by_username
            ]

            if missing_users:
                raise ValueError(
                    "These mapped usernames do not exist in the "
                    "local database: "
                    + ", ".join(missing_users)
                )

            users_by_sheet_name = {
                sheet_name: users_by_username[username]
                for sheet_name, username in USER_MAPPING.items()
            }

            # Load existing submissions for duplicate detection.
            existing_songs = session.exec(select(Song)).all()
            songs_by_key = {}

            for song in existing_songs:
                key = song_key(
                    song.title,
                    song.artist,
                    song.submission_date,
                    song.submitter_id,
                )
                songs_by_key.setdefault(key, song)

            # Track existing votes so rerunning the script is safe.
            existing_votes = session.exec(select(Vote)).all()
            votes_by_key = {
                (vote.song_id, vote.voter_id)
                for vote in existing_votes
            }

            songs_added = 0
            votes_added = 0
            rows_skipped = 0
            invalid_ratings = 0

            # Spreadsheet layout:
            # D = date, E = poster, I = song title, J = artist.
            # K:R = ratings from the eight members.
            last_date = None
            last_poster = None

            # Row 1 contains the headers.
            for row_number, row in enumerate(
                sheet.iter_rows(min_row=2, values_only=True),
                start=2,
            ):
                # D (index 3) is the song-level date.
                # A (index 0) is the fallback date.
                row_date = parse_date(row[3]) if len(row) > 3 and row[3] else None

                if row_date is None and row and row[0]:
                    row_date = parse_date(row[0])

                if row_date is not None:
                    last_date = row_date

                # E (index 4) identifies the poster/submitter.
                poster = clean_text(row[4]) if len(row) > 4 else ""

                if poster:
                    last_poster = poster

                # I and J contain the song title and artist.
                title = clean_text(row[8]) if len(row) > 8 else ""
                artist = clean_text(row[9]) if len(row) > 9 else ""

                # Ignore blank spreadsheet rows.
                if not title and not artist:
                    continue

                if not title or not artist:
                    raise ValueError(
                        f"Row {row_number}: song title or artist "
                        "is missing."
                    )

                if last_date is None:
                    raise ValueError(
                        f"Row {row_number}: no usable submission date."
                    )

                if not last_poster:
                    raise ValueError(
                        f"Row {row_number}: no poster/submitter found."
                    )

                # Match the spreadsheet name without case sensitivity.
                poster_key = last_poster.strip().casefold()
                mapped_poster = next(
                    (
                        name
                        for name in USER_MAPPING
                        if name.casefold() == poster_key
                    ),
                    None,
                )

                if mapped_poster is None:
                    raise ValueError(
                        f"Row {row_number}: poster '{last_poster}' "
                        "is not in USER_MAPPING."
                    )

                submitter = users_by_sheet_name[mapped_poster]

                key = song_key(
                    title,
                    artist,
                    last_date,
                    submitter.id,
                )

                song = songs_by_key.get(key)

                if song is None:
                    song = Song(
                        title=title,
                        artist=artist,
                        submitter_id=submitter.id,
                        submission_date=last_date,
                    )
                    session.add(song)
                    session.flush()

                    songs_by_key[key] = song
                    songs_added += 1
                else:
                    rows_skipped += 1

                # Import each member's rating as a vote.
                for member_name, column_number in RATING_COLUMNS.items():
                    cell_index = column_number - 1
                    rating_value = (
                        row[cell_index]
                        if len(row) > cell_index
                        else None
                    )

                    if rating_value is None or not clean_text(rating_value):
                        continue

                    score = parse_rating(rating_value)

                    if score is None:
                        invalid_ratings += 1
                        print(
                            f"Warning: row {row_number}, "
                            f"{member_name}: ignored invalid rating "
                            f"{rating_value!r}"
                        )
                        continue

                    voter = users_by_sheet_name[member_name]
                    vote_key = (song.id, voter.id)

                    if vote_key in votes_by_key:
                        continue

                    session.add(
                        Vote(
                            song_id=song.id,
                            voter_id=voter.id,
                            score=score,
                        )
                    )

                    votes_by_key.add(vote_key)
                    votes_added += 1

            # Commit only after processing the worksheet successfully.
            session.commit()

            print("\nHistorical import completed.")
            print(f"Local database: {DATABASE_FILE}")
            print(f"New songs imported: {songs_added}")
            print(f"New votes imported: {votes_added}")
            print(f"Existing song submissions reused: {rows_skipped}")
            print(f"Invalid ratings ignored: {invalid_ratings}")
            print("\nNo averages, winners, or rankings were imported.")

    finally:
        workbook.close()


if __name__ == "__main__":
    main()