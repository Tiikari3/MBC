from pathlib import Path
import sqlite3
from openpyxl import load_workbook

# The database beside this script
BACKEND_DIR = Path(__file__).resolve().parent
DB_PATH = BACKEND_DIR / "mbc_local.db"


def main():
    workbook_path = input(
        "Enter the full path to your Excel workbook: "
    ).strip().strip('"')

    excel_path = Path(workbook_path).expanduser()

    if not excel_path.is_file():
        print(f"\nERROR: File not found: {excel_path}")
        return

    if excel_path.suffix.lower() != ".xlsx":
        print("\nERROR: Please select an .xlsx file.")
        return

    workbook = load_workbook(
        excel_path,
        read_only=True,
        data_only=True,
    )

    print("\n=== EXCEL WORKBOOK ===")
    print(f"File: {excel_path}")

    for sheet in workbook.worksheets:
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows, ())

        # Ignore completely empty header cells at the end.
        headers = list(headers)
        while headers and headers[-1] is None:
            headers.pop()

        data_rows = sum(
            1 for row in rows
            if any(value is not None for value in row)
        )

        print(f"\nSheet: {sheet.title}")
        print(f"Data rows: {data_rows}")
        print("Column headers:")

        for index, header in enumerate(headers, start=1):
            print(f"  {index}. {header!r}")

    workbook.close()

    print("\n=== SQLITE DATABASE ===")
    print(f"Database path: {DB_PATH}")
    print(f"Database exists: {DB_PATH.exists()}")

    if not DB_PATH.exists():
        print("No database found at this path.")
        return

    with sqlite3.connect(DB_PATH) as connection:
        tables = connection.execute("""
            SELECT name
            FROM sqlite_master
            WHERE type = 'table'
              AND name NOT LIKE 'sqlite_%'
            ORDER BY name
        """).fetchall()

        if not tables:
            print("No tables found.")
            return

        for (table_name,) in tables:
            # Table names come from SQLite's own table listing.
            safe_name = table_name.replace('"', '""')
            columns = connection.execute(
                f'PRAGMA table_info("{safe_name}")'
            ).fetchall()

            count = connection.execute(
                f'SELECT COUNT(*) FROM "{safe_name}"'
            ).fetchone()[0]

            print(f"\nTable: {table_name} ({count} records)")
            print("Columns:", ", ".join(
                column[1] for column in columns
            ))


if __name__ == "__main__":
    main()