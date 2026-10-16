from pathlib import Path
from openpyxl import load_workbook

# Ask for the Excel spreadsheet's location
excel_path = input(
    "Enter the full path to your Excel file: "
).strip().strip('"')

excel_path = Path(excel_path)

if not excel_path.exists():
    print(f"File not found: {excel_path}")
    raise SystemExit(1)

# Read the spreadsheet without changing it
workbook = load_workbook(
    excel_path,
    read_only=False,
    data_only=True
)

for sheet in workbook.worksheets:
    print("\n" + "=" * 70)
    print(f"SHEET: {sheet.title}")
    print(f"Dimensions: {sheet.max_row} rows x {sheet.max_column} columns")

    merged = list(sheet.merged_cells.ranges)
    if merged:
        print("Merged cells:", merged[:20])

    print("\nFirst 12 rows (non-empty cells only):")

    for row_number in range(1, min(12, sheet.max_row) + 1):
        values = []

        for column_number in range(1, sheet.max_column + 1):
            cell = sheet.cell(row_number, column_number)

            if cell.value is not None:
                values.append(
                    f"{cell.coordinate}={cell.value!r}"
                )

        if values:
            print(f"Row {row_number}: " + " | ".join(values))

workbook.close()
print("\nInspection complete. No spreadsheet data was changed.")